import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
print("slide size", p.slide_width, p.slide_height, "slides", len(p.slides))
print("layouts:", [l.name for l in p.slide_layouts])
for i, s in enumerate(p.slides, 1):
    print(f"\n=== slide {i} layout={s.slide_layout.name}")
    for sh in s.shapes:
        kind = sh.shape_type
        t = ""
        if sh.has_text_frame:
            t = " | ".join(par.text for par in sh.text_frame.paragraphs if par.text.strip())
        extra = ""
        if sh.has_table if hasattr(sh, "has_table") else False:
            tb = sh.table
            extra = " TABLE " + " || ".join(" | ".join(c.text for c in r.cells) for r in tb.rows)
        if getattr(sh, "has_chart", False) and sh.has_chart:
            extra = " CHART"
        print(f"  [{kind}] {sh.name} pos=({Emu(sh.left).inches if sh.left is not None else None:.2f},{Emu(sh.top).inches if sh.top is not None else 0:.2f}) size=({Emu(sh.width).inches if sh.width else 0:.2f}x{Emu(sh.height).inches if sh.height else 0:.2f}) {t[:300]}{extra[:600]}")
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        print("  NOTES:", s.notes_slide.notes_text_frame.text[:300])
