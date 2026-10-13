import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
for i, s in enumerate(p.slides, 1):
    texts = [sh.text_frame.text.replace("\n"," / ") for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    pics = sum(1 for sh in s.shapes if sh.shape_type == 13)
    tabs = sum(1 for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table)
    charts = sum(1 for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart)
    print(f"{i:2d} [{len(list(s.shapes))} shapes, pic {pics}, tab {tabs}, chart {charts}] {texts[0][:60] if texts else ''} :: {texts[1][:110] if len(texts)>1 else ''}")
