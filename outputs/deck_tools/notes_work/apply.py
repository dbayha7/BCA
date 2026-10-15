import copy, json, sys
from pptx import Presentation
from pptx.util import Emu
W = sys.argv[1]; DECK = sys.argv[2]; OUT = sys.argv[3]
JOURNAL = "/root/.claude/projects/-home-user/9e341792-006d-5a11-a8e9-8fac16ac7f5b/subagents/workflows/wf_42d93be2-4a8/journal.jsonl"
old = json.load(open(f"{W}/notes_work/slides.json"))
def key(texts):
    return (texts[0] if texts else "", texts[1] if len(texts) > 1 else "")
oldkey = {o["slide"]: key(o["text"]) for o in old}
# writer results from the journal
by_key = {}
for line in open(JOURNAL):
    o = json.loads(line)
    if o.get("type") != "result":
        continue
    r = o.get("result") or {}
    if isinstance(r, str):
        try: r = json.loads(r)
        except Exception: continue
    for n in (r.get("notes") or []):
        k = oldkey.get(n["slide"])
        if k: by_key[k] = n["notes"]
mine = json.load(open(f"{W}/notes_work/mine.json"))
p = Presentation(DECK)
body = None
for s in p.slides:
    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
        body = s.notes_slide.notes_placeholder._element; break
done, kept, missing = 0, 0, []
for i, s in enumerate(p.slides, 1):
    shapes = sorted(s.shapes, key=lambda sh: (round(Emu(sh.top).inches, 1) if sh.top is not None else 0,
                                              Emu(sh.left).inches if sh.left is not None else 0))
    texts = [sh.text_frame.text.strip() for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    k = key(texts)
    text = mine.get(k[1]) or by_key.get(k)
    if text is None:
        cands = [v for kk, v in by_key.items() if kk[0] == k[0] and k[0]]
        if len(cands) == 1 and not k[1]:
            text = cands[0]
    if text is None:
        kept += 1; missing.append((i, k[1][:50])); continue
    ns = s.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(copy.deepcopy(body))
    ns.notes_text_frame.text = text
    done += 1
p.save(OUT)
print("new notes", done, "| kept old", kept)
for m in missing: print("  kept:", m)
