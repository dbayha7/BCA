import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
want = sys.argv[2].split(',')
for spec in want:
    i, name = spec.split(':')
    s = p.slides[int(i)-1]
    for sh in s.shapes:
        if sh.name == name and sh.has_text_frame:
            for pi, par in enumerate(sh.text_frame.paragraphs):
                for r in par.runs:
                    c=None
                    try: c=r.font.color.rgb
                    except: pass
                    print(i, name, pi, repr(r.text[:80]), r.font.size.pt if r.font.size else None, r.font.bold, c)
