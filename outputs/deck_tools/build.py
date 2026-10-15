"""Revise the WBCP deck: delete the old toy-world slides, add studies 6-8 (ThrusterChain, why BCA's channel cannot
reach the harm, the new channel), and update the slides that refer to the study list."""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu

SRC, OUT, CH = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(SRC)
S = list(prs.slides)                      # original slides, S[n - 1] is slide n
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def rgb(h):
    return RGBColor.from_string(h)


def shp(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(f"{name} not on slide")


def set_runs(shape, texts, para=0):
    """Replace the run texts of one paragraph, keeping each run's formatting; extra runs are removed."""
    p = shape.text_frame.paragraphs[para]
    runs = p.runs
    for i, t in enumerate(texts):
        runs[i].text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def set_text(slide, name, text, para=0):
    set_runs(shp(slide, name), [text], para)


def drop_paras(shape, keep):
    for p in shape.text_frame.paragraphs[keep:]:
        p._p.getparent().remove(p._p)


def delete(slide, *names):
    for n in names:
        el = shp(slide, n)._element
        el.getparent().remove(el)


def move(slide, names, dy=0.0, dx=0.0):
    for n in names:
        s = shp(slide, n)
        s.top = s.top + Inches(dy)
        s.left = s.left + Inches(dx)


def fill(slide, name, fill_hex, text_hex=None, line_hex=None):
    s = shp(slide, name)
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill_hex)
    if line_hex:
        s.line.color.rgb = rgb(line_hex)
    if text_hex:
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = rgb(text_hex)


def clone(src):
    new = prs.slides.add_slide(src.slide_layout)
    for s in list(new.shapes):
        s._element.getparent().remove(s._element)
    bg = src._element.cSld.find(P_NS + "bg")
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for el in src.shapes._spTree:
        tag = etree.QName(el).localname
        if tag in ("nvGrpSpPr", "grpSpPr", "pic"):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    return new


NOTE_BODY = None
for _sl in S:
    if _sl.has_notes_slide and _sl.notes_slide.notes_text_frame is not None:
        NOTE_BODY = _sl.notes_slide.notes_placeholder._element
        break


def notes(slide, text):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(copy.deepcopy(NOTE_BODY))
    tf = ns.notes_text_frame
    tf.text = text


PILL = {"Met": ("2F6DB5", "FFFFFF"), "Missed": ("D9822B", "FFFFFF"), "Partly met": ("A9C4E6", "14213D"),
        "Pending": ("E3E1DA", "4A5568"), "Built": ("2F6DB5", "FFFFFF"), "Bar set": ("A9C4E6", "14213D"),
        "Exact": ("2F6DB5", "FFFFFF"), "Method": ("14213D", "FFFFFF"), "Controls": ("2F6DB5", "FFFFFF")}
STATUS = {"FINISHED": ("2F6DB5", "FFFFFF"), "RUNNING": ("F0B070", "14213D"), "STOPPED": ("D9822B", "FFFFFF")}


def intro(eyebrow, title, why, does, setup, status, status_text, plain, source, running_template=False):
    s = clone(S[49])                       # slide 50: a finished study's opening slide
    set_text(s, "Text 0", eyebrow); set_text(s, "Text 1", title)
    set_text(s, "Text 4", why); set_text(s, "Text 7", does)
    for name, t in zip(("Text 11", "Text 13", "Text 15", "Text 17"), setup):
        set_text(s, name, t)
    set_text(s, "Status 201", status); fill(s, "Status 201", *STATUS[status])
    set_text(s, "Text 202", status_text); set_text(s, "Text 205", plain)
    set_runs(shp(s, "Source 301"), ["Source  ", source])
    return s


def chart(eyebrow, title, change, why, img, exp, got, pill, plain, exp_label=None, got_label=None):
    s = clone(S[56])                       # slide 57: a change slide with one chart
    set_text(s, "Text 0", eyebrow); set_text(s, "Text 1", title)
    set_runs(shp(s, "Text 2"), ["CHANGE", "  " + change])
    set_runs(shp(s, "Text 3"), ["WHY IT MAKES SENSE", "  " + why])
    s.shapes.add_picture(f"{CH}/{img}.png", Inches(1.07), Inches(1.74), Inches(11.19), Inches(4.31))
    set_text(s, "Text 6", exp); set_text(s, "Text 9", got)
    if exp_label:
        set_text(s, "Text 5", exp_label)
    if got_label:
        set_text(s, "Text 8", got_label)
    set_text(s, "Text 10", pill); fill(s, "Text 10", PILL[pill][0], PILL[pill][1], PILL[pill][0])
    set_runs(shp(s, "Plain terms 201"), ["IN PLAIN TERMS  ", plain])
    return s


def cards(eyebrow, title, items, change, why, exp, got, pill, plain, change_label="CHANGE",
          why_label="WHY IT MAKES SENSE", exp_label=None, got_label=None):
    s = clone(S[55])                       # slide 56: three cards
    set_text(s, "Text 0", eyebrow); set_text(s, "Text 1", title)
    for (lab, head, formula, desc), names in zip(items, (("Text 3", "Text 4", "Text 5", "Text 6"),
                                                         ("Text 8", "Text 9", "Text 10", "Text 11"),
                                                         ("Text 13", "Text 14", "Text 15", "Text 16"))):
        for n, t in zip(names, (lab, head, formula, desc)):
            set_text(s, n, t)
        shp(s, names[3]).height = Inches(1.35)
    for n in ("Shape 2", "Shape 7", "Shape 12"):
        shp(s, n).height = Inches(3.12)
    move(s, ["Text 17", "Text 18", "Shape 19", "Text 20", "Text 21", "Shape 22", "Text 23", "Text 24", "Text 25"],
         dy=0.86)
    set_runs(shp(s, "Text 17"), [change_label, "  " + change])
    if why:
        set_runs(shp(s, "Text 18"), [why_label, "  " + why])
    else:
        delete(s, "Text 18"); shp(s, "Text 17").width = Inches(11.5)
    set_text(s, "Text 21", exp); set_text(s, "Text 24", got)
    if exp_label:
        set_text(s, "Text 20", exp_label)
    if got_label:
        set_text(s, "Text 23", got_label)
    set_text(s, "Text 25", pill); fill(s, "Text 25", PILL[pill][0], PILL[pill][1], PILL[pill][0])
    # an IN PLAIN TERMS line, styled like the chart slides' line
    tb = s.shapes.add_textbox(Inches(0.89), Inches(5.86), Inches(11.8), Inches(0.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = "IN PLAIN TERMS  "
    r1.font.size = Pt(11.5); r1.font.bold = True; r1.font.name = "Arial"; r1.font.color.rgb = rgb("2F6DB5")
    r2 = p.add_run(); r2.text = plain
    r2.font.size = Pt(12.5); r2.font.name = "Arial"; r2.font.color.rgb = rgb("14213D")
    return s


def table_slide(eyebrow, title, tables, plain, note_lines=()):
    s = clone(S[56])
    set_text(s, "Text 0", eyebrow); set_text(s, "Text 1", title)
    delete(s, "Text 2", "Text 3", "Shape 4", "Text 5", "Text 6", "Shape 7", "Text 8", "Text 9", "Text 10")
    for (left, top, width, head, colw, rows, font, rowh) in tables:
        if head:
            tb = s.shapes.add_textbox(Inches(left), Inches(top - 0.38), Inches(width), Inches(0.3))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = head
            r.font.size = Pt(13); r.font.bold = True; r.font.name = "Arial"; r.font.color.rgb = rgb("2F6DB5")
        gf = s.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width),
                                Inches(rowh * len(rows)))
        tbl = gf.table
        tbl.first_row = True; tbl.horz_banding = False
        for j, w in enumerate(colw):
            tbl.columns[j].width = Inches(w)
        for i, row in enumerate(rows):
            tbl.rows[i].height = Inches(rowh)
            for j, val in enumerate(row):
                c = tbl.cell(i, j); c.text = val
                c.fill.solid(); c.fill.fore_color.rgb = rgb("14213D" if i == 0 else ("FDFCF9" if i % 2 else "F1EFE8"))
                c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(font); r.font.name = "Arial"; r.font.bold = (i == 0)
                        r.font.color.rgb = rgb("FFFFFF" if i == 0 else "14213D")
    y = 6.2 - 0.3 * max(0, len(note_lines) - 2)
    for ln in note_lines:
        tb = s.shapes.add_textbox(Inches(0.89), Inches(y), Inches(11.8), Inches(0.3))
        r = tb.text_frame.paragraphs[0].add_run(); r.text = ln
        r.font.size = Pt(12); r.font.name = "Courier New"; r.font.color.rgb = rgb("14213D")
        y += 0.3
    set_runs(shp(s, "Plain terms 201"), ["IN PLAIN TERMS  ", plain])
    return s


def adds(eyebrow, title, counts, why_not, means, nxt):
    s = clone(S[53])                       # slide 54: what a study adds up to
    set_text(s, "Text 0", eyebrow); set_text(s, "Text 1", title)
    set_text(s, "Text 11", why_not); set_text(s, "Text 14", means); set_text(s, "Text 17", nxt)
    shp(s, "Text 17").width = Inches(10.0)
    if counts is None:
        delete(s, "Text 2", "Text 3", "Text 4", "Text 5", "Text 6", "Text 7", "Text 8")
        move(s, ["Shape 9", "Text 10", "Text 11", "Shape 12", "Text 13", "Text 14", "Shape 15", "Text 16", "Text 17"],
             dy=-0.62)
        return s
    (lead, segs) = counts                   # segs: [(n, label), ...] for met, missed, other
    total = sum(n for n, _ in segs)
    x, full, gap = 0.89, 11.56, 0.02
    usable = full - gap * (sum(1 for n, _ in segs if n) - 1)
    for (n, lab), name in zip(segs, ("Text 2", "Text 3", "Text 4")):
        sh = shp(s, name)
        if n == 0:
            sh._element.getparent().remove(sh._element)
            continue
        w = usable * n / total
        sh.left, sh.width = Inches(x), Inches(w)
        set_runs(sh, [str(n)])
        x += w + gap
    set_text(s, "Text 5", lead)
    lx = 0.89 + 0.09 * len(lead) + 0.12
    shp(s, "Text 5").width = Inches(0.09 * len(lead) + 0.1)
    for (n, lab), name in zip(segs, ("Text 6", "Text 7", "Text 8")):
        sh = shp(s, name)
        if n == 0:
            sh._element.getparent().remove(sh._element)
            continue
        set_runs(sh, [lab])
        w = 0.095 * len(lab) + 0.12
        sh.left, sh.width = Inches(lx), Inches(w)
        lx += w + 0.08
    return s


# ------------------------------------------------------------------ study 6: ThrusterChain and its harm checks
new = []
s = intro("STUDY 6 OF 8 · A CONTROLLED SHIFT",
          "When missing data hide a trap, is there room for BCA's caution?",
          "Toy worlds could not show whether caution lands where the critic is wrong.",
          "Builds a world with exact truth and a known gap, and checks the harm before judging BCA.",
          ["ThrusterChain: exact values", "Densest 15% of rows removed", "4 situations, 3 control worlds",
           "Pre-registered gate; 5 seeds"],
          "STOPPED",
          "Run to its pre-registered stop on 12 Oct: the harm is real, but there is no room for BCA's kind of caution. "
          "BCA was not run, as designed.",
          "We built a simulated channel where pushing a thruster past 80% backfires, then removed almost every "
          "past-80% command from the training data. The critic trained on that data thinks full throttle is fine, and "
          "the learned controller pays for it when the current is stronger than anything recorded. Could extra "
          "caution of BCA's kind win that back?",
          "outputs/shift/stop-report/README.md · experiments/shift/SHIFT_ENV.md · pre-registration 68d9e877…")
notes(s, "Study 6 replaces the toy world. The 3-state toy world could not test BCA: no thin data, planted errors, a "
         "6-number policy. Its replacement, the Gapped LQ world, failed its own gate twice (next slide). So we built "
         "ThrusterChain, where truth is exact by simulation and the data gap is created on purpose. The gate checks "
         "the harm first, and only lets BCA be judged if caution of BCA's kind has room to help.")
new.append(s)

s = chart("STUDY 6 · WHY A NEW WORLD: THE TOY WORLD'S REPLACEMENT",
          "The Gapped LQ world failed its own gate twice",
          "a 10-state world with thin, rare contexts", "the 3-state toy world had no thin data", "c1_glq",
          "Loss sits in thin data in ≥ 4 of 5 seeds", "At most 3 of 5, at every alpha, in both gates",
          "Missed",
          "Its learned critic was wrong almost everywhere (best action off by 1.8–4.8 against noise 0.3), so it "
          "could not test placement.")
notes(s, "Gate R3 asks whether the return lost sits where the data are thin. First gate (seeds 0-4): base world 0/1/0 "
         "of 5 seeds at alpha 2.5/5/10; the thin-share retune reached 3 of 5 at best; retune 2 was unstable. "
         "Amendment A1 (post hoc, labelled, fresh seeds 5-9): base 0/3/2, thin-share 0/0/0. The critic is inaccurate "
         "even where data are plentiful, so the world cannot isolate harm from missing data. BCA was never run there.")
new.append(s)

s = chart("STUDY 6 · THE ENVIRONMENT: THRUSTERCHAIN",
          "Five floats, four thrusters, and a stall past 80%",
          "a Gymnasium environment with a hidden trap", "every value is exact, by simulation", "c2_env",
          "Densest 15% removed: 99% of full-throttle rows gone",
          "Current 0.7–1.1 (recorded 0.1–0.7); big waves; big kicks",
          "Built",
          "The recording controller's most common command is full throttle; after the removal the data show almost "
          "nothing past 80%.",
          exp_label="THE GAP", got_label="THE SHIFT")
notes(s, "The task: hold five floats at rest against a current and waves, using four thrusters between neighbouring "
         "floats. Observation: 13 numbers (each float's position and speed, the current, two wave numbers). Action: "
         "4 throttles in [-1, 1]. Reward each step: minus the squared distance from rest minus 0.3 times the squared "
         "throttles; 100 steps. The trap: past 80% a thruster stalls; its push falls from 0.8 to 0.5 at full "
         "throttle, while its energy cost keeps rising. A no-stall comparison world is identical except push = "
         "throttle everywhere, so there the missing data hide nothing. Data: 1,000 episodes x 100 steps from a "
         "recording controller; the D4RL densest-15% rule removes 99% of the full-throttle commands. Evaluation: "
         "recorded conditions, strong current (0.7-1.1, beyond anything recorded), big waves, big starting kicks.")
new.append(s)

s = chart("STUDY 6 · CHECKS 1–2: IS THE HARM REAL?",
          "The missing data cost 99 points in strong current",
          "same actor, full log vs gapped log", "only the missing data differ",
          "c3_harm",
          "Critic ≥ 2× as wrong at removed commands; harm ≥ 5",
          "2.15–2.40× as wrong; harm 98.9 [96.2, 101.7]",
          "Met",
          "Random removal (3.1 points) and the no-stall world (0.4) explain almost none of it: the missing stall data "
          "cause the harm.")
notes(s, "Check 1 passed 6 of 6 marks: at the removed commands the gapped critic is 2.15-2.40x as wrong as the "
         "full-log critic; far from them 0.68-0.80x; random removal 0.96-1.21x; another initialisation 0.87-1.09x. "
         "Pushing past 80% truly costs 0.33-0.35; the gapped critic overrates it by 0.33-0.36, so it sees no cost. "
         "Check 2 passed at alpha 5 (one mark missed at 2.5). Points: 0 = recording controller, 100 = stall-aware "
         "controller. Plain TD3+BC on the gapped log beats cloning by 27-47 points in three situations but loses 71 "
         "in strong current.")
new.append(s)

s = chart("STUDY 6 · CHECK 3: IS THERE ROOM?",
          "Caution of BCA's kind cannot win it back",
          "extra pull, placed by the exact harm", "an upper bound for BCA's dose",
          "c4_room",
          "Harm-placed beats matched controls by ≥ 3 points",
          "−1.9 to +1.1 points; no situation certified",
          "Missed",
          "Only even ×3 caution moves strong current (+17 of 99 points), and it costs 5–13 points everywhere "
          "else. The gate stopped before BCA.")
notes(s, "This is the key negative result. Each arm is an actor trained against the gapped critic with extra pull "
         "toward the recorded command on each training row, exactly BCA's form (the behaviour-cloning term times a "
         "per-row factor in [1, 1.5)). The harm-placed arm uses the exact harm of each row, better information than "
         "BCA could ever learn. It beats strength-matched even caution, amount-matched and shuffled controls by -1.9 "
         "to +1.1 points; the mark was 3. Factors up to 2 do no better. So the limit is the channel, not the signal.")
new.append(s)

s = adds("STUDY 6 · WHAT IT ADDS UP TO", "The harm is real, but BCA's channel cannot reach it",
         ("4 checks:", [(2, "2 passed"), (1, "1 failed"), (1, "1 not reached")]),
         "The pull acts on logged rows, toward logged commands. In strong current the actor is far from both.",
         "A better signal alone cannot help here. The channel the signal acts through has to change.",
         "Study 7 · Why can't BCA's caution reach this harm?")
new.append(s)

# ------------------------------------------------------------------ study 7: why the channel fails
s = intro("STUDY 7 OF 8 · WHY THE CHANNEL FAILS",
          "Why can't BCA's caution reach harm that lies outside the data?",
          "Study 6 found no room. Before changing BCA, we needed the exact reason.",
          "Traces the harm to one wrong belief, and shows where BCA reads its signal and acts.",
          ["BCA's code, read line by line", "Study 6's saved critics, actors", "Exact truth by simulation",
           "No BCA run; 5 seeds"],
          "FINISHED",
          "Diagnosis complete on 13 Oct. Exploratory: computed outside the gate on its saved artifacts; no "
          "pre-registered rule was changed.",
          "BCA looks for trouble at the rows in the data, and its fix is to pull the actor back toward the logged "
          "command. But the critic is right at those rows. It is wrong only about commands the data never show, past "
          "80%. The actor learns that wrong belief at ordinary training states and carries it into strong current.",
          "outputs/shift/beyond-the-data/mechanism.md, anatomy.md · algorithms/td3_bc_bca.py · "
          "calibration/dose.py")
new.append(s)

s = cards("STUDY 7 · THREE REASONS, FROM BCA'S CODE",
          "BCA reads and acts where the critic is right",
          [("1", "Reads at logged rows", "dose(sᵢ, aᵢ)",
            "The band is fitted, calibrated and read only at logged (state, command) rows. There the gapped critic "
            "is accurate (error 0.07). It is never read at the actor's own command."),
           ("2", "Pulls toward the recorder", "mᵢ · (π(sᵢ) − aᵢ)²",
            "The dose only scales the pull toward the logged command (throttle 1 ≈ 0.5). That says ‘drive "
            "like the recorder’, not ‘stop at 80%’, so it cannot cancel the critic's false slope."),
           ("3", "Capped, and blind to bias", "1 ≤ dose < 1.5",
            "The dose is capped at 1.5, and ‘cannot certify’ gives dose 1: no caution. The error BCA "
            "calibrates cannot see a critic that is consistently wrong, since it matches its own targets.")],
          "read BCA's code from fit to dose", "where is the signal read and used?",
          "Caution where the critic is wrong", "Caution where the critic is right", "Missed",
          "The signal and the fix both live at the data; the error lives at the actor's own commands past 80%.")
notes(s, "From the code: BCA's scale network is trained on one-step errors at recorded training rows; the threshold "
         "is calibrated on held-out recorded rows; the dose is read at the training batch's recorded rows "
         "(algorithms/td3_bc_bca.py). Nothing evaluates the width at the actor's command pi(s). In TD3+BC the dose "
         "multiplies the behaviour-cloning pull on that recorded row (algorithms/td3_bc.py), so a dose of m acts like "
         "using alpha/m on that row, between alpha/1.5 and alpha. An infinite width (cannot certify) gives dose 1 "
         "(calibration/dose.py). And the one-step error is computed against targets built from the same critic, so a "
         "critic that consistently overrates pushing past 80% leaves small errors at recorded rows.")
new.append(s)

s = chart("STUDY 7 · EVIDENCE 1: THE CRITIC'S BELIEF",
          "The critic is right up to 80% and wrong past it",
          "compare each critic with the exact values", "find which belief the actor learns", "c5_belief",
          "The gapped critic errs where data are missing", "Error +0.51 past 80%; 0.07 inside the data",
          "Met",
          "At training states where the water pushes hardest, the gapped critic sees no cost past 80%; critics that "
          "saw the stall data do.")
notes(s, "Values of thruster 1's throttle relative to holding it at 80%, averaged over training states in the "
         "full-throttle region (water's push above 0.565), other thrusters from the stall-aware controller. Truth is "
         "exact. The gapped critic matches the truth to within 0.01 up to 80% and then keeps rising. Its error in the "
         "preference for full throttle over 80% is +0.51 [0.48, 0.53]; the full-log critic -0.02, random removal "
         "-0.00. Inside the recorded range the gapped critic is no worse than the others (mean absolute error 0.07 "
         "against 0.12), which is exactly where BCA looks.")
new.append(s)

s = chart("STUDY 7 · EVIDENCE 2: ONE BELIEF CAUSES THE HARM",
          "Cap the actor at 80% and it scores 101.8 instead of −14.3",
          "cap only the past-80% commands, when the actor runs", "isolates the one wrong belief", "c6_onebelief",
          "The harm comes from commands past 80%", "−14.3 → +101.8 points in strong current",
          "Met",
          "Everything else the gapped critic taught is good: the fix needed is one value, at the actor's own commands, "
          "at training states.")
notes(s, "The actor is trained only at training states, and every actor update reads the critic at the actor's own "
         "command there. The gapped critic's actor pushes past 80% at 42% of training states in the full-throttle "
         "region, against 10% for the full-log critic's actor. As the current grows beyond the recorded range its "
         "commands keep rising, so the same mistake is made at most steps in strong current. Cutting only those "
         "commands back to 80% at evaluation turns -14.3 into +101.8 points, above the full-log critic's actor "
         "(84.6).")
new.append(s)

s = adds("STUDY 7 · WHAT IT ADDS UP TO", "BCA acts where the critic is right", None,
         "Changing only where the band is read is not enough for TD3+BC: study 6 put the exact harm in this channel "
         "and found no room.",
         "Read the band at the actor's own command, act on the value the actor climbs, drop the 1.5 cap, and treat "
         "‘cannot certify’ as most caution.",
         "Study 8 · Does a new channel work?")
new.append(s)

# ------------------------------------------------------------------ study 8: the new channel
s = intro("STUDY 8 OF 8 · A NEW CHANNEL",
          "Can caution reach the harm if it acts on the value the actor climbs?",
          "Study 7 found the error sits at the actor's own commands, not at the data.",
          "Tests the new channel with perfect information, sets the bar, then tests BCA in it.",
          ["Stage 1: exact error, no BCA", "Data-only rival signals", "Trace test bed (harm claim)",
           "Stage 2: BCA, pre-registered"],
          "RUNNING",
          "Stage 1, the rivals and the trace test bed are finished. Stage 2: BCA's plain band failed detection; with "
          "classifier weights it passed. The policy test on seeds 0–4 is running.",
          "Instead of pulling the actor toward the data, lower the critic's value at the actor's own command by how "
          "much it is overrated there. With perfect information this wins back almost all of the harm, and it is "
          "cautious only where the missing commands are bad.",
          "outputs/shift/beyond-the-data/PLAN.md, stage1/RESULTS.md, stage2/PREREGISTRATION.md (v6, 9b38c7c4…)")
new.append(s)

s = chart("STUDY 8 · STAGE 1: THE NEW CHANNEL, PERFECT INFORMATION",
          "The new channel wins back 94 of 99 points",
          "penalty = exact error at the actor's command", "an upper bound for any signal in this slot", "c7_stage1",
          "≥ 20 points over matched controls; recover ≥ 49.5", "+89 to +99 over controls; recovered 93.9",
          "Met",
          "Untargeted caution of equal strength recovers nothing, and the right information at the wrong state "
          "recovers little: placement is the effect.")
notes(s, "Exploratory, no BCA. Spec and code committed before any return was computed; controls matched in strength "
         "blind and locked first. Every actor starts from the cloned actor and trains 20,000 steps against study 6's "
         "frozen gapped critic. Penalty: how much more the critic overrates the actor's command than the row's "
         "recorded command, by the exact one-step error. Marks (fixed in advance) all pass: beats each control by "
         "92.4 / 89.4 / 99.0 / 94.2 points, 5 of 5 seeds; recovers 93.9 (half the harm was the mark); keeps the "
         "improvement elsewhere (recorded +7.9 over no caution, big waves +19.6, big kicks +16.9).")
new.append(s)

s = chart("STUDY 8 · RIVALS: THE BAR FOR BCA",
          "A perfect signal is cautious only where caution is needed",
          "data-only signals in the same slot", "the bar BCA has to clear", "c8_rivals",
          "Some data-only signal captures the room", "Only the range hinge: +84.6, but it costs 20.4",
          "Bar set",
          "Clipping to the recorded range works only because every missing command is bad here; where they help, it "
          "costs 20.")
notes(s, "Each rival is a data-only signal used in the same lower-edge slot, matched in strength. Gain: points won "
         "back over no caution in strong current in the stalling world. Cost: points lost in strong current in the "
         "no-stall world, where the missing commands are good. Only the range hinge captures the room (84.6). It is "
         "the rival Stage 2 compares BCA against. The perfect signal's no-stall cost (1.3) is from Stage 1.")
new.append(s)

s = chart("STUDY 8 · TRACE TEST BED: CAN THE HARM ITSELF BE SEEN?",
          "A trace of the stall in the data stayed invisible",
          "keep 0.5–10% of the full-throttle rows", "to test whether a signal finds the harm itself", "c9_trace",
          "A share with harm ≥ 50 and separation ≥ 0.2", "Separation ≤ 0.08 at every share",
          "Missed",
          "Seeds 0–2; shares 1–10% were added after the first look (post hoc). Testing the harm claim needs a "
          "new design.")
notes(s, "The honest limit of study 6's world: there every missing command is bad, so a signal that only detects "
         "novelty looks as good as a perfect one. To test whether BCA identifies the harm itself (not just "
         "unsupported commands), the data must carry a trace of the harm. We kept a small share of the removed "
         "full-throttle rows. Harm stays above 50 points up to 5%, but the data-only residual signal never separates "
         "the stalling world from the no-stall world (at most 0.08 against a mark of 0.2). No share works.")
new.append(s)

s = cards("STUDY 8 · STAGE 2: BCA IN THE NEW CHANNEL",
          "BCA's band, read at the actor's own command",
          [("1", "Read at the actor's command", "U(s, π(s))",
            "BCA's calibrated width, evaluated at the actor's own command and paired against the logged command's "
            "width, so only the extra uncertainty counts."),
           ("2", "Act on the value climbed", "Q − k·[U(π) − U(a)]₊",
            "The actor climbs the critic's lower edge, with k = 1. No 1.5 cap; ‘cannot certify’ now means "
            "the most caution."),
           ("3", "Check when it runs", "flag → fall back",
            "The same band screens each command while the controller runs; a flagged command falls back toward the "
            "recorded range.")],
          "pre-registered v6, approved 13 Oct", "the flags claim, confirmed on fresh seeds",
          "BCA flags unsupported commands better than rivals", "Plain band: missed. Weighted: 3 of 4 tests", "Partly met",
          "Confirmatory: fresh seeds 5–9 with every setting frozen, about 36–46 hours on 2 processes.")
notes(s, "The flags claim: BCA's band, read at the actor's own command, flags commands the data do not support, and "
         "used as a penalty it recovers return in strong current at a small cost in the no-stall world. It is judged "
         "against the best data-only rival (the range hinge) and matched untargeted controls, on gate seeds 0-4 "
         "(exploratory) and then fresh seeds 5-9 (confirmatory). The claim sentence fixed in advance: in strong "
         "current BCA's lower edge recovers X of H points and keeps Y of TD3+BC's gain; the range limit recovers Z; "
         "in the no-stall world BCA costs W and the range limit V.")
new.append(s)

s = adds("STUDY 8 · WHAT IT ADDS UP TO", "BCA flags the problem with weights; policy test next",
         ("8 checks so far:", [(5, "5 met"), (2, "2 missed"), (1, "1 pending: policy test")]),
         "BCA's plain band barely changes past 80%. With weights it works, but no better than the classifier alone "
         "(0.86 vs 0.90 in strong current).",
         "Whether BCA's penalty wins back return is now the question. The policy test on seeds 0–4 answers it.",
         "Policy test, deployment check and full host, seeds 0–4: running")
new.append(s)

# ------------------------------------------------------------------ the environment, step by step (Farama order), and the methods
(intro6, glq, env, checks12, check3, adds6, intro7, cards7, belief, onebelief, adds7,
 intro8, stage1, rivals, trace, stage2, adds8) = new

# 1 of 7: description (the diagram slide)
set_text(env, "Text 0", "STUDY 6 · THRUSTERCHAIN 1 OF 7 · DESCRIPTION")
set_text(env, "Text 1", "The task: hold five floats at rest against the water")
set_runs(shp(env, "Text 2"), ["WHAT", "  a Gymnasium environment, BCAShift/ThrusterChain-v0"])
set_runs(shp(env, "Text 3"), ["WHY THIS WORLD", "  every value is exact, by simulation"])
set_text(env, "Text 5", "THE GOAL"); set_text(env, "Text 6", "Floats at rest, using as little throttle as possible")
set_text(env, "Text 8", "THE TWIST"); set_text(env, "Text 9", "Past 80% a thruster stalls: push falls to 0.5 at full")
set_runs(shp(env, "Plain terms 201"), ["IN PLAIN TERMS  ", "A good controller pushes up to 80% and no further. The data will "
                                       "hide that fact, and the test is whether a learner finds out the hard way."])
notes(env, "Five floats sit in a row in a channel, tied together by weak springs. A steady current and waves push "
           "every float downstream. Four thrusters sit between neighbouring floats: thruster 1 pushes floats 1 and 2, "
           "thruster 2 pushes floats 2 and 3, and so on. The controller holds the floats at rest while spending little "
           "throttle. The twist: past 80% throttle the propeller stalls and the push falls (0.65 at 90%, 0.5 at full "
           "throttle, against 0.8 at 80%), while the throttle still costs energy. The next slides go through the "
           "standard Gymnasium card: spaces, dynamics, reward, start, episode end, versions; then the data and the "
           "reference controllers; then the methods; then the results.")

spaces = table_slide(
    "STUDY 6 · THRUSTERCHAIN 2 OF 7 · ACTION AND OBSERVATION SPACES",
    "Four throttles in, thirteen numbers out",
    [(0.89, 1.75, 5.3, "ACTION SPACE · Box(−1, 1, (4,), float32)", [0.9, 3.2, 1.2],
      [["Index", "Action", "Range"],
       ["0", "throttle of thruster 1 (floats 1 and 2)", "−1 to 1"],
       ["1", "throttle of thruster 2 (floats 2 and 3)", "−1 to 1"],
       ["2", "throttle of thruster 3 (floats 3 and 4)", "−1 to 1"],
       ["3", "throttle of thruster 4 (floats 4 and 5)", "−1 to 1"]], 15, 0.55),
     (6.45, 1.75, 6.0, "OBSERVATION SPACE · Box(low, high, (13,), float32)", [1.35, 3.25, 1.4],
      [["Index", "Observation", "Range"],
       ["0, 2, 4, 6, 8", "position of floats 1–5 (0 = at rest)", "−300 to 300"],
       ["1, 3, 5, 7, 9", "speed of floats 1–5", "−300 to 300"],
       ["10", "strength of the current, c", "0 to 1.5"],
       ["11, 12", "the wave's place in its 23-step cycle, as two clock hands (w₁, w₂)", "−1 to 1"]], 15, 0.62)],
    "step() rounds each command to float32 and clips it to [−1, 1]. The ±300 bound is proven: from any allowed start "
    "the floats stay within 256.4 of rest.",
    note_lines=("import experiments.shift.thruster_chain   # registers the ids",
                "env = gymnasium.make(\"BCAShift/ThrusterChain-v0\")   # check_env passes"))

dyn = cards("STUDY 6 · THRUSTERCHAIN 3 OF 7 · TRANSITION DYNAMICS",
            "Floats swing, water pushes, thrusters push back",
            [("1", "The floats swing", "A · s",
              "Each float is a damped swing (periods 63, 31, 18, 11 and 8 steps; it loses 2–4% per step). Weak "
              "springs (size 0.02) pass a little of each float's motion to the others."),
             ("2", "The thrusters push", "0.3 · B · thrust(a)",
              "Thruster k pushes the speeds of floats k and k+1. Push = throttle up to 80%; past it "
              "0.8 − 1.5(|a| − 0.8): 0.65 at 90%, 0.5 at full throttle."),
             ("3", "The water pushes", "0.3 · v · (c + w₁)",
              "The same push on every float: a current c that is fixed for the episode, plus a wave w₁ that turns "
              "0.27 radians per step (a 23.3-step cycle).")],
            "s′ = A·s + 0.3·B·thrust(a) + 0.3·v·(c + w₁);  c′ = c;  the wave turns",
            "", "Identical, except push = throttle all the way",
            "Below 80% the worlds match; data can't tell them apart", "Exact",
            "s: the ten float numbers. a: the four throttles. No noise anywhere, so every value is exact by simulation.",
            change_label="NEXT STATE", exp_label="NO-STALL WORLD", got_label="WHY IT MATTERS")
notes(dyn, "One deterministic dynamics function, in numpy and in JAX with the same arithmetic order; step(), the data "
           "generator and the exact truth all call it. A is the five-oscillator plant (radii 0.98-0.96, angles 0.10 to "
           "0.80 radians) plus a seeded random coupling of size 0.02, kept small so the world stays stable without "
           "thrust (largest growth factor 0.98) while the floats are not five independent problems. The no-stall "
           "comparison world is identical except thrust(a) = a, so there the removed commands are harmless.")

reward = cards("STUDY 6 · THRUSTERCHAIN 4 OF 7 · REWARD, STARTING STATE, EPISODE END",
               "What is scored, how an episode starts, and when it ends",
               [("1", "Reward", "r = −(|s|² + 0.3·|a|²)",
                 "Every step: the floats' squared distance from rest plus 0.3 × the squared throttles. Energy is "
                 "paid even while a thruster stalls."),
                ("2", "Starting state", "s₀ ~ N(0, 0.5²)",
                 "Float positions and speeds drawn N(0, 0.5²); current 0.1–0.7; wave size 0.1–0.4; the wave can be "
                 "anywhere in its cycle."),
                ("3", "Episode end", "t = 100",
                 "Truncation only, never terminated. A return is the 100-step sum; learners and the truth use "
                 "discount 0.98.")],
               "100 rewards summed; closer to 0 is better",
               "0 = recording controller, 100 = stall-aware controller",
               "1,000 fixed starts per situation, the same for everyone", "Every controller is simulated, not estimated",
               "Exact",
               "points = 100 × (return − recording controller's) ÷ (stall-aware controller's − recording controller's), "
               "per situation.",
               change_label="RETURN", why_label="POINTS", exp_label="A FIXED TEST", got_label="EXACT SCORES")

versions = table_slide(
    "STUDY 6 · THRUSTERCHAIN 5 OF 7 · VERSIONS AND ARGUMENTS",
    "Eight registered versions: two worlds × four situations",
    [(0.89, 1.75, 11.55, "", [3.9, 1.6, 3.85, 2.2],
      [["Id", "World", "Start law (what changes)", "Outside the recorded range"],
       ["BCAShift/ThrusterChain-v0", "stalling", "recorded conditions: current 0.1–0.7, wave 0.1–0.4, float spread 0.5",
        "none (training conditions)"],
       ["BCAShift/ThrusterChain-StrongCurrent-v0", "stalling", "current 0.7–1.1", "every start"],
       ["BCAShift/ThrusterChain-BigWaves-v0", "stalling", "wave size 0.4–0.7", "every start"],
       ["BCAShift/ThrusterChain-BigKicks-v0", "stalling", "float start spread 1.5", "on part of the steps"]],
      14, 0.66)],
    "Strong current and big waves lie entirely outside the recorded conditions: no recorded episode had a current "
    "above 0.7 or a wave above 0.4.",
    note_lines=("BCAShift/ThrusterChainNoStall-v0 (+ the same 3 variants): the no-stall world, same four start laws",
                "ThrusterChainEnv(stall=True, start=\"recorded\", render_mode=None, stall_slope=1.5)",
                "info per step: state, t, push, thrust, stalled, command   ·   render_mode=\"ansi\": one line of text"))

data = cards("STUDY 6 · THRUSTERCHAIN 6 OF 7 · THE DATA AND THE GAP",
             "The data: one recording controller, then a deliberate gap",
             [("1", "Recording controller", "87% normal driving",
               "Holds against the water, wobbles (noise 0.3), never past 80%. Where the water pushes hardest, half "
               "the time it gives full throttle on thruster 1: the log's most common command."),
              ("2", "The log", "1,000 × 100 steps",
               "Split by episode: 500 training (50,000 rows), 200 calibration (a 1,000-row bank), 300 held out. D4RL "
               "format; every row replays exactly."),
              ("3", "The gap", "densest 15% removed",
               "The D4RL rule removes the commands in the densest ball: 99.1–99.8% of full-throttle commands. Kept "
               "rows past 80%: 0.04–0.14%, from 12–14%.")],
             "full log · the same number of rows removed at random · the no-stall world",
             "", "46–48% of those states stay, with normal commands", "In the commands, not in the states",
             "Controls",
             "The training data still visit the states where full throttle matters, but almost never show what full "
             "throttle does there.",
             change_label="CONTROLS", exp_label="WHAT STAYS", got_label="SO THE HOLE IS")

refctl = table_slide(
    "STUDY 6 · THRUSTERCHAIN 7 OF 7 · REFERENCE CONTROLLERS",
    "Fixed reference controllers set the points scale",
    [(0.89, 1.75, 11.55, "", [2.5, 3.45, 1.4, 1.4, 1.4, 1.4],
      [["Controller", "What it does", "Recorded", "Strong current", "Big waves", "Big kicks"],
       ["recording controller", "the log's source, both modes", "−79.8 (0)", "−253.9 (0)", "−106.8 (0)", "−263.2 (0)"],
       ["stall-aware controller", "textbook controller, capped at 80%", "−26.4 (100)", "−103.8 (100)", "−36.2 (100)",
        "−163.5 (100)"],
       ["normal driving only", "the normal mode, without wobble", "−37.9 (78.4)", "−146.9 (71.3)", "−55.6 (72.5)",
        "−205.4 (58.0)"],
       ["rising-thrust controller", "textbook controller, capped at 100%: assumes thrust keeps rising", "−35.1 (83.6)",
        "−335.2 (−54.1)", "−61.8 (63.8)", "−219.2 (44.2)"]], 14, 0.72)],
    "Exact 100-step returns (points in brackets), stalling world. The stall-aware controller is a reference, not "
    "proven optimal, so scores above 100 are possible.")
notes(refctl, "The textbook controller is the discounted linear-quadratic regulator of the model in which thrust keeps "
              "rising with throttle. Capped at 80% it is the stall-aware controller (100 points). Capped at 100% it is "
              "the rising-thrust controller, which is right in the no-stall world and badly wrong in strong current "
              "in the stalling world (-54 points): exactly the mistake a critic that never saw the stall would make.")

m1 = cards("STUDY 6 · METHOD 1 OF 3 · THE LEARNER (TD3+BC)",
           "The learner: a critic rates, an actor picks",
           [("1", "Critic Q(s, a)", "y = r + 0.98·Q(s′, c(s′))",
             "Rates a command at a state, trained only on logged rows to match one-step targets. In the checks it "
             "scores following the stall-aware controller c, so its error is exact."),
            ("2", "Actor π(s)", "λ·Q1(s, π(s)) − (π(s) − a)²",
             "Picks the command. It climbs the critic's value at its own command, while a behaviour-cloning term "
             "pulls it toward the logged command a."),
            ("3", "Alpha", "λ = α / mean|Q|",
             "How far the actor may trust the critic over the data. A ladder fixed in advance: 2.5, then 5, then "
             "10, the same for every method.")],
           "one critic trained on the full log, one on the gapped log",
           "same start, rows and batches for both actors",
           "Pure imitation of the gapped log: every actor's start", "Only what the critic learned from the missing rows",
           "Method",
           "TD3+BC (Fujimoto & Gu, 2021) is the learner BCA sits on. Plain TD3+BC, critic and actor trained together, "
           "is also reported.",
           change_label="TWO CRITICS", why_label="FAIR", exp_label="CLONED ACTOR", got_label="THE ONLY DIFFERENCE")

m2 = cards("STUDY 6 · METHOD 2 OF 3 · BCA AS IT IS TODAY",
           "BCA today: a calibrated band, turned into extra pull",
           [("1", "Band", "U = R · u · η(s, a)",
             "η: a small network predicting the size of the critic's one-step error, fitted at logged rows. R: one "
             "threshold from held-out rows (WBCP), so about 90% of errors fall inside."),
            ("2", "Dose", "m = 1 + 0.5·U/(U + u)",
             "The band's width becomes a multiplier between 1 and 1.5. A band too wide to certify gives m = 1."),
            ("3", "Where it acts", "m · (π(s) − a)²",
             "The dose multiplies the behaviour-cloning pull on each logged row: a wider band means a harder pull "
             "toward the logged command.")],
           "the same BCA as studies 1–5", "can extra pull on logged rows reach this harm?",
           "Wide band where the critic is wrong → caution there", "Check 3 gives this channel the exact harm",
           "Method",
           "u is the typical error size (a running average). Nothing in this channel ever looks at the actor's own "
           "command.",
           change_label="RECAP", why_label="THE QUESTION", exp_label="THE HOPE", got_label="HOW WE TEST IT")

m3 = table_slide(
    "STUDY 6 · METHOD 3 OF 3 · THE HARM CHECKS",
    "Four checks, fixed in advance, before BCA may be judged",
    [(0.89, 1.75, 11.55, "", [0.75, 2.75, 4.65, 3.4],
      [["Check", "Question", "Pass mark (pre-registered)", "Rules out"],
       ["1", "Is the critic wrong where data are missing?",
        "Gapped critic ≥ 2× as wrong as the full-log critic at the removed commands; about 1× far from them",
        "random removal; another initialisation; a world without the stall"],
       ["2", "Does acting on that error cost return?",
        "Harm ≥ 5 points, ≥ 3 more than random removal and the no-stall world explain",
        "having less data; harm that needs no stall"],
       ["3", "Could caution of BCA's kind help? (room)",
        "Pull placed by the exact harm, factors in [1, 1.5), beats strength-matched even, same-total and shuffled "
        "pull by ≥ 3 points",
        "caution that helps by its amount, not its placement"],
       ["4", "Is there power to see BCA's effect?", "Run only for situations check 3 certifies", "—"]], 13.5, 0.8)],
    "Pre-registration 68d9e877… (commit bdb8c8a) hashed every file. BCA's band is fitted only if checks 1–3 pass. "
    "Seeds 0–4.")
notes(m3, "Check 3 is the key design choice. It asks whether caution in BCA's own form, the behaviour-cloning term "
          "multiplied by a per-row factor in BCA's range, can help at all if it is placed by the exact harm of each "
          "row: information BCA could never have. Controls are matched blind before any return is computed: even "
          "pull of the same strength (same change to the actor), the same total amount spread evenly, and the same "
          "factors shuffled across rows. If perfect placement cannot beat them, no BCA signal can.")

m4 = cards("STUDY 8 · METHOD · THE NEW CHANNEL",
           "Lower the value where the critic overrates the actor's own command",
           [("1", "Lower edge", "Q1(s, π(s)) − k · penalty",
             "The actor climbs the critic's value minus a penalty at its own command, at training states. There is no "
             "pull toward the data."),
            ("2", "Stage 1 penalty", "[err(s,π) − err(s,a)]₊",
             "err: the critic's exact one-step error, by simulation. The penalty is how much more it overrates the "
             "actor's command than the logged one."),
            ("3", "Controls", "matched strength",
             "Even extra pull, push less everywhere, two random penalty surfaces: each tuned blind to change the actor "
             "as much, and locked before any return.")],
           "if perfect information fails here, no BCA version can", "≥ 20 over every control; recover ≥ 49.5",
           "Seeds 0–4; every actor starts from the cloned actor", "Stalling world; the no-stall world gives the cost",
           "Method",
           "Stage 2 puts BCA's band in place of the exact error; that is the version under test now.",
           change_label="WHY EXACT FIRST", why_label="MARKS", exp_label="SETUP", got_label="WORLDS")

m5 = table_slide(
    "STUDY 8 · METHOD · RIVAL SIGNALS AND THE TRACE TEST",
    "The rivals BCA must beat, and how the trace test works",
    [(0.89, 1.75, 11.55, "", [2.9, 8.65],
      [["Signal (data only)", "What it measures at a state and command"],
       ["range hinge", "how far each throttle goes past the recorded range (squared)"],
       ["novelty: density, distance", "how far the command is from the recording controller's usual command, or from the "
        "nearest training rows"],
       ["classifier ratio", "a classifier's odds that the command came from the actor rather than the log"],
       ["two-head disagreement", "the gap between the critic's two copies, Q1 and Q2"]], 15, 0.62)],
    "Trace test: keep some removed full-throttle rows; pass if harm stays ≥ 50 and a signal (residual evidence: the "
    "nearby rows' one-step error) flags past-80% commands ≥ 0.2 more often in the stalling world.")

m6 = table_slide(
    "STUDY 8 · METHOD · HOW DETECTION IS SCORED",
    "How BCA's detection is scored: four tests",
    [(0.89, 1.75, 11.55, "", [1.9, 6.15, 3.5],
      [["Test", "Question", "Pass mark (pooled, and in ≥ 4 of 5 seeds)"],
       ["Direction", "Is the band wider at the actor's harmful past-80% command than at that command capped at 80%?",
        "≥ 75% of pairs (training states); ≥ 65% (strong current)"],
       ["Size", "Is the extra penalty big enough to cancel the critic's false preference, with k = 1?",
        "median ratio ≥ 1"],
       ["Ranking", "Does a harmful command score above a harmless one?",
        "≥ 0.75 (training states); ≥ 0.70 (strong current)"],
       ["Beat the rivals", "Is BCA's ranking within 0.02 of the best data-only signal?", "on both sets of states"]],
      14, 0.78)],
    "Pairs: 1,500 training states where the water pushes hardest, 1,500 other training states, 3,000 strong-current "
    "states. Seeds 0–4 are exploratory; only fresh seeds 5–9 can carry a claim.")
notes(m6, "Harmful means the command really is worse than its 80% cap, by exact simulation. Calibration is also checked: "
          "about 5% of held-out logged rows should lie above BCA's cut-off. If direction or size fails with uniform "
          "weights, the pre-registration declares the weighted version (likelihood-ratio weights in WBCP) and repeats "
          "the test. If that also fails, no BCA training arm runs; the deployment check and the report remain.")
m6r = chart("STUDY 8 · STAGE 2 FIRST RESULT: DETECTION, UNIFORM VERSION",
            "BCA's band barely changes past 80%",
            "BCA's band at the actor's command, seeds 0–4", "the first look, fixed in advance", "c10_detect",
            "Direction, size, ranking pass; within 0.02 of rivals",
            "Size 0.0004 (mark 1); direction 0.50; below rivals",
            "Missed",
            "Calibration holds (4.9% above the cut-off; target 5%). As pre-registered, the weighted version now runs; "
            "if it also fails, no BCA training arm runs.")
notes(m6r, "Uniform version, seeds 0-4, finished 13 Oct 13:22 UTC. Direction: the band is wider at the harmful command "
           "than at its 80% cap in 50% of training-state pairs (mark 75%; 1 of 5 seeds) and 73% of strong-current pairs "
           "(mark 65%, but only 2 of 5 seeds). Size: the median paired penalty is 0.0004 of the critic's false "
           "preference (mark 1): the band does not widen where it must. Ranking: 0.80 on training states (passes) and "
           "0.62 in strong current (mark 0.70). The classifier ratio, a data-only signal, ranks 0.99 and 0.90. The "
           "claim rule is not met; seeds 0-4 could not carry a claim anyway.")
m6w = chart("STUDY 8 · STAGE 2 SECOND RESULT: DETECTION, WEIGHTED VERSION",
            "With classifier weights, BCA's band flags the problem",
            "add likelihood-ratio weights from a classifier", "the declared fallback, fixed in advance", "c11_detect_both",
            "Direction, size, ranking pass; within 0.02 of rivals",
            "All three pass; 0.86 vs 0.90 in strong current",
            "Partly met",
            "The weights come from the same classifier as the best rival, so the detection mostly reflects that "
            "classifier. The policy test now runs.")
new = [intro6, glq, env, spaces, dyn, reward, versions, data, refctl, m1, m2, m3, checks12, check3, adds6,
       intro7, cards7, belief, onebelief, adds7, intro8, m4, stage1, m5, rivals, trace, stage2, m6, m6r, m6w, adds8]

# ------------------------------------------------------------------ edits to existing slides
# title slide
t1 = S[0]
set_text(t1, "Text 3", "Eight studies. For each: why, what changed, what we expected, what we got. Studies 6–8 are new.")
set_text(t1, "Text 4", "David Bayha · October 13, 2026")
if t1.has_notes_slide:
    tf = t1.notes_slide.notes_text_frame
    tf.text = tf.text.replace("There are seven studies", "There are eight studies").replace("seven studies", "eight studies")

# "STUDY N OF 7" -> "OF 8"
for n in (20, 25, 31, 39, 45):
    sh = shp(S[n - 1], "Text 0")
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.text = r.text.replace("OF 7", "OF 8")

# setup slide 7: the second test bed is now ThrusterChain
set_text(S[6], "Text 29", "Studies 6–8 use a second test bed: ThrusterChain, a simulated task where the true values "
                          "are known exactly and the data gap is made on purpose.")

# slide 10: drop the two toy-problem rows, lift the last row and the notes below it
s10 = S[9]
rm, lift = [], []
for sh in list(s10.shapes):
    top = Emu(sh.top).inches
    if 3.37 <= top < 5.06:
        rm.append(sh)
    elif top >= 5.06:
        lift.append(sh)
for sh in rm:
    sh._element.getparent().remove(sh._element)
for sh in lift:
    sh.top = sh.top - Inches(1.72)
sh91 = shp(s10, "Text 91")
set_runs(sh91, [sh91.text_frame.paragraphs[0].runs[0].text.replace("Studies 6–7 train the actor on a built critic",
                                                                      "Studies 6–8 use a simulated task")])
p2 = sh91.text_frame.paragraphs[1]
set_runs(sh91, ["The true result is known exactly, and the data gap is made on purpose."], para=1)

# slide 16, 17: drop references to the deleted studies
sh = shp(S[15], "Text 109")
set_runs(sh, [sh.text_frame.paragraphs[0].runs[0].text.replace(" Studies 5 and 6 come back to this.", " Study 5 comes back to this.")])
sh = shp(S[16], "Text 8")
for p in sh.text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace(" → studies 6–7", " → studies 6–8").replace("studies 6–7", "studies 6–8")

# roadmap (slide 19): cards 6 and 7 become the new studies 6 and 7; card 8 is added; 8 cards re-laid out
rd = S[18]
set_text(rd, "Text 1", "Eight studies; each answer raises the next question")
card = {k: [f"Shape {2 + 5 * (k - 1)}", f"Text {3 + 5 * (k - 1)}", f"Text {4 + 5 * (k - 1)}",
            f"Text {5 + 5 * (k - 1)}", f"Text {6 + 5 * (k - 1)}", f"Finding {200 + k}"] for k in range(1, 8)}
set_text(rd, "Text 29", "Test world"); set_text(rd, "Text 30", "Is there room to help?")
set_text(rd, "Text 31", "FINISHED")
set_text(rd, "Finding 206", "99 points lost; no room for BCA's channel", para=1)
set_text(rd, "Text 33", "7"); set_text(rd, "Text 34", "The channel"); set_text(rd, "Text 35", "Why can't caution reach it?")
set_text(rd, "Text 36", "FINISHED")
for n in ("Text 33", "Text 36"):
    for p in shp(rd, n).text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = rgb("2F6DB5")
shp(rd, "Text 36").top = shp(rd, "Text 31").top
set_text(rd, "Finding 207", "Read where the critic is right", para=1)
# card 8: copy card 7's shapes
src_names = card[7]
copies = []
for n in src_names:
    el = copy.deepcopy(shp(rd, n)._element)
    rd.shapes._spTree.append(el)
    copies.append(rd.shapes[-1])
c8 = {"card": copies[0], "num": copies[1], "title": copies[2], "q": copies[3], "status": copies[4], "find": copies[5]}
set_runs(c8["num"], ["8"]); set_runs(c8["title"], ["New channel"]); set_runs(c8["q"], ["Does a new channel work?"])
set_runs(c8["status"], ["RUNNING NOW"]); set_runs(c8["find"], ["Works with weights; policy test running"], para=1)
for n in ("num", "status"):
    for p in c8[n].text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = rgb("A3530E")
c8["status"].top = Inches(4.15); c8["status"].height = Inches(0.49)
step, cw, tw = 1.457, 1.36, 1.22
for k in range(1, 9):
    if k <= 7:
        objs = [shp(rd, n) for n in card[k]]
    else:
        objs = copies
    x0 = 0.89 + (k - 1) * step
    objs[0].left, objs[0].width = Inches(x0), Inches(cw)
    for o in objs[1:]:
        o.left, o.width = Inches(x0 + 0.08), Inches(tw)
split = 0.89 + 3 * step + cw
shp(rd, "Shape 37").width = Inches(split - 0.89)
shp(rd, "Shape 40").left = Inches(split + 0.097); shp(rd, "Shape 40").width = Inches(12.45 - split - 0.097)
for n in ("Text 41", "Text 42"):
    shp(rd, n).left = Inches(split + 0.097); shp(rd, n).width = Inches(12.45 - split - 0.097)
set_text(rd, "Text 42", "studies 5–8 check what it does to learning")

# study 5's last slide points to study 6
set_text(S[48], "Text 19", "Study 6 · When missing data hide a trap, is there room for BCA's caution?")

# summary (slide 60): rows 6-7 replaced, row 8 added
sm = S[59]
set_text(sm, "Text 1", "Eight questions, eight answers")
set_text(sm, "Text 29", "Is there room in a controlled shift?"); set_text(sm, "Text 31", "Harm of 99 points; no room for BCA's channel")
set_text(sm, "Text 34", "Why can't caution reach it?"); set_text(sm, "Text 36", "Read and applied where the critic is right")
for p in shp(sm, "Text 33").text_frame.paragraphs:
    for r in p.runs:
        r.font.color.rgb = rgb("2F6DB5")
row8 = []
for n in ("Shape 32", "Text 33", "Text 34", "Text 35", "Text 36"):
    el = copy.deepcopy(shp(sm, n)._element)
    sm.shapes._spTree.append(el)
    o = sm.shapes[-1]; o.top = o.top + Inches(0.48); row8.append(o)
set_runs(row8[1], ["8"]); set_runs(row8[2], ["Does a new channel work?"])
set_runs(row8[4], ["Yes with perfect information; BCA needs classifier weights"])
for p in row8[1].text_frame.paragraphs:
    for r in p.runs:
        r.font.color.rgb = rgb("D9822B")

# NEXT (slide 61)
nx = S[60]
items = [("Policy test, seeds 0–4", "weighted BCA's penalty against the range hinge and matched controls: running"),
         ("Deployment, full host", "the same band as a run-time check, and inside full TD3+BC training"),
         ("Fresh seeds 5–9", "the confirmatory run, every setting frozen"),
         ("Harm claim, real data", "a trace test bed that works; the D4RL action-gap test")]
for k, (head, body) in enumerate(items):
    set_text(nx, f"Text {4 + 4 * k}", head); set_text(nx, f"Text {5 + 4 * k}", body)
delete(nx, "Shape 18", "Text 19", "Text 20", "Text 21")
for k in range(4):
    x = 0.89 + k * 2.92
    shp(nx, f"Shape {2 + 4 * k}").left = Inches(x); shp(nx, f"Shape {2 + 4 * k}").width = Inches(2.8)
    for t in (3, 4, 5):
        o = shp(nx, f"Text {t + 4 * k}"); o.left = Inches(x + 0.17); o.width = Inches(2.45)

# appendix index (slide 63)
sh = shp(S[62], "Text 3")
for p in sh.text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace(", alignment and placement code", " and alignment code")

# ------------------------------------------------------------------ at most four items per slide
def gloss_split(src, groups, first, total):
    out = []
    for gi, ks in enumerate(groups):
        g = clone(src)
        keep = set()
        for k in ks:
            keep |= {f"Term {k}", f"Text {k + 1}", f"Text {k + 2}"}
        for sh in list(g.shapes):
            nm = sh.name
            is_term = nm.startswith("Term ") or (nm.startswith("Text ") and nm not in ("Text 0", "Text 1", "Text 3"))
            if is_term and nm not in keep:
                sh._element.getparent().remove(sh._element)
        for j, k in enumerate(ks):
            col, row = j % 2, j // 2
            x, y = 0.89 + col * 6.05, 1.6 + row * 2.4
            card, ttl, desc = shp(g, f"Term {k}"), shp(g, f"Text {k + 1}"), shp(g, f"Text {k + 2}")
            card.left, card.top, card.width, card.height = Inches(x), Inches(y), Inches(5.75), Inches(2.15)
            ttl.left, ttl.top, ttl.width, ttl.height = Inches(x + 0.22), Inches(y + 0.18), Inches(5.35), Inches(0.42)
            desc.left, desc.top, desc.width, desc.height = Inches(x + 0.22), Inches(y + 0.66), Inches(5.35), Inches(1.4)
            for p in ttl.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(17)
            for p in desc.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(14.5)
        e = shp(g, "Text 0")
        for p in e.text_frame.paragraphs:
            for r in p.runs:
                r.text = r.text.replace("(1 OF 2)", f"({first + gi} OF {total})").replace("(2 OF 2)", f"({first + gi} OF {total})")
        out.append(g)
    return out


gl_a = gloss_split(S[3], [[4, 7, 10], [13, 16, 19], [22, 25, 28]], 1, 6)        # the toy-problem term (31) is dropped
gl_b = gloss_split(S[4], [[4, 7, 10, 13], [16, 19, 22, 25], [28, 34, 37]], 4, 6)  # "comparison doses" (31) is dropped


def region_split(src, keep_left, card_top=(1.7, 4.9)):
    """Roadmap: keep the four cards on one side (by position) and spread them across the slide."""
    g = clone(src)
    lo, hi = card_top
    step_old, x_old = 1.457, 0.89
    for sh in list(g.shapes):
        t = Emu(sh.top).inches
        if not (lo <= t < hi):
            continue
        k = int(round((Emu(sh.left).inches - x_old) / step_old))
        if (k < 4) != keep_left:
            sh._element.getparent().remove(sh._element)
            continue
        j = k if keep_left else k - 4
        off = Emu(sh.left).inches - (x_old + k * step_old)
        sh.left = Inches(0.89 + j * 2.92 + off)
        sh.width = Inches(2.8 if off < 0.04 else 2.62)
    return g


road_a = region_split(rd, True)
road_b = region_split(rd, False)
for g, which in ((road_a, 1), (road_b, 2)):
    set_text(g, "Text 0", f"ROADMAP ({which} OF 2)")
delete(road_a, "Shape 40", "Text 41", "Text 42")
shp(road_a, "Shape 37").width = Inches(11.56)
for n in ("Text 38", "Text 39"):
    shp(road_a, n).width = Inches(11.56)
set_text(road_a, "Text 1", "Studies 1–4: is BCA's guarantee real?")
delete(road_b, "Shape 37", "Text 38", "Text 39")
for n in ("Shape 40", "Text 41", "Text 42"):
    shp(road_b, n).left = Inches(0.89); shp(road_b, n).width = Inches(11.56)
set_text(road_b, "Text 1", "Studies 5–8: does it help the actor?")


def rows_split(src, first_half):
    g = clone(src)
    cut = 1.55 + 0.48 * 4
    for sh in list(g.shapes):
        t = Emu(sh.top).inches
        if t < 1.5:
            continue
        if (t < cut) != first_half:
            sh._element.getparent().remove(sh._element)
        elif not first_half:
            sh.top = sh.top - Inches(0.48 * 4)
    return g


sum_a = rows_split(sm, True); sum_b = rows_split(sm, False)
set_text(sum_a, "Text 1", "Eight questions, eight answers (1 of 2)")
set_text(sum_b, "Text 1", "Eight questions, eight answers (2 of 2)")

REPLACE = {4: gl_a, 5: gl_b, 19: [road_a, road_b], 60: [sum_a, sum_b]}

# ------------------------------------------------------------------ tables of at most four rows
def cell_set(cell, text):
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = text
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def table_of(slide):
    for sh in slide.shapes:
        if sh.has_table:
            return sh
    raise KeyError("no table")


# "Types of Sampling": merge the two K-rows-per-episode samplers into one row
ts = table_of(S[5]).table
r3, r4 = list(ts.rows)[3], list(ts.rows)[4]
cell_set(r3.cells[0], "K rows per episode (random or evenly spaced)")
cell_set(r3.cells[1], "About n/K episodes, with K rows from each: at random positions, or one from each of K equal "
                      "stretches of the episode")
r4._tr.getparent().remove(r4._tr)

# appendix "Where do the write-ups disagree?": two slides of four rows
def rows_keep(src, keep, label):
    g = clone(src)
    gf = table_of(g)
    rows = list(gf.table.rows)
    for i, r in enumerate(rows):
        if i > 0 and i not in keep:
            r._tr.getparent().remove(r._tr)
    t = shp(g, "Text 1")
    set_runs(t, [t.text_frame.paragraphs[0].runs[0].text + label])
    return g


dis_a = rows_keep(S[73], {1, 2, 3, 4}, " (1 of 2)")
dis_b = rows_keep(S[73], {5, 6, 7, 8}, " (2 of 2)")
REPLACE[74] = [dis_a, dis_b]

# ------------------------------------------------------------------ scorecard (after the roadmap)
sc = clone(nx)                                     # the NEXT slide, already cut to four cards
set_text(sc, "Text 0", "SCORECARD")
set_text(sc, "Text 1", "The guarantee is sounder; there is no win for learning yet")
cards4 = [("Studies 1–5", "Calibration: flaws found, fixed",
           "Four real flaws in how BCA sets its cut-off: wrong method, blind to shift, related rows in its bank, the "
           "wrong critic copy. Each one now has a fix, though it isn't installed in BCA everywhere yet."),
          ("No study yet", "BCA has not improved a policy",
           "Studies 1–5 used frozen critics, so none measured learning. The only return number is one seed on one "
           "dataset, and it is descriptive only."),
          ("Studies 6–7", "Today's channel can't reach the harm",
           "BCA adds extra pull toward the logged data. Even placed perfectly, that pull wins back almost nothing of "
           "the 99 points the data gap costs."),
          ("Study 8", "A channel that works; BCA's band, with help",
           "A value penalty at the actor's own command wins back 94 of 99 points with perfect information. BCA's "
           "plain band failed; with classifier weights it flags the problem. Policy test running.")]
for k, (num, head, body) in enumerate(cards4):
    set_text(sc, f"Text {3 + 4 * k}", num)
    set_text(sc, f"Text {4 + 4 * k}", head)
    set_text(sc, f"Text {5 + 4 * k}", body)
    o = shp(sc, f"Text {3 + 4 * k}")
    for p in o.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(15)
    shp(sc, f"Text {4 + 4 * k}").height = Inches(0.6)
    b = shp(sc, f"Text {5 + 4 * k}"); b.top = Inches(3.05); b.height = Inches(1.6)
    shp(sc, f"Text {4 + 4 * k}").top = Inches(2.35)
    for p in b.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12.5)
set_text(sc, "Text 22", "In one sentence: studies 1–5 made the guarantee trustworthy; studies 6–8 asked whether it helps "
                        "learning, and so far it does not in its current form.")
shp(sc, "Text 22").width = Inches(11.56)
REPLACE[19] = [road_a, road_b, sc]

# ------------------------------------------------------------------ order: delete toy-world slides, insert the new ones
DELETE = {8, 46, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59, 73}
order = []
for n in range(1, len(S) + 1):
    if n in DELETE:
        continue
    if n in REPLACE:
        order.extend(REPLACE[n])
    else:
        order.append(S[n - 1])
    if n == 49:
        order.extend(new)
sldIdLst = prs.slides._sldIdLst
by_slide = {}
for sid in list(sldIdLst):
    by_slide[prs.part.related_part(sid.rId)] = sid
for sid in list(sldIdLst):
    sldIdLst.remove(sid)
keep_parts = set()
for sl in order:
    sldIdLst.append(by_slide[sl.part]); keep_parts.add(sl.part)
for part, sid in by_slide.items():
    if part not in keep_parts:
        prs.part.drop_rel(sid.rId)
prs.save(OUT)
print("slides", len(prs.slides), "new", len(new))
