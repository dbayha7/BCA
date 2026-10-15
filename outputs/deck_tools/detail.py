import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
for idx in map(int, sys.argv[2].split(',')):
    s = p.slides[idx-1]
    print(f"\n######## slide {idx}  bg={s.background.fill.type if s.background else None}")
    for sh in s.shapes:
        d = f"[{sh.shape_type}] {sh.name!r} ({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}"
        try:
            if sh.fill.type is not None:
                d += f" fill={sh.fill.fore_color.rgb if sh.fill.type==1 else sh.fill.type}"
        except Exception as e: pass
        try:
            if sh.line.fill.type == 1: d += f" line={sh.line.color.rgb} w={sh.line.width}"
        except Exception: pass
        print(d)
        if sh.has_text_frame:
            for par in sh.text_frame.paragraphs:
                runs = par.runs
                if not runs: continue
                r = runs[0]; f = r.font
                col = None
                try: col = f.color.rgb
                except Exception: pass
                print(f"     p: size={f.size.pt if f.size else None} bold={f.bold} font={f.name} color={col} align={par.alignment} :: {par.text[:150]}")
        if sh.shape_type == 13:
            print("     picture", sh.image.content_type, sh.image.size)
