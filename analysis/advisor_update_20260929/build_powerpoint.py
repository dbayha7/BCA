import json, hashlib, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

BASE=Path(__file__).parent
REPO=BASE.parents[1]
SRC=REPO/'outputs/advisor_interpretation/2026-09-29/BCA_Results_and_Weighted_CP.pptx'
OUT=REPO/'outputs/advisor_update/2026-09-29-clear-story'
OUT.mkdir(parents=True,exist_ok=True)
data=json.loads((BASE/'authored-content.json').read_text(encoding='utf-8'))
native=json.loads((BASE/'raw-template.json').read_text(encoding='utf-8'))
p=Presentation(SRC)
assert len(p.slides)==len(native['slides'])==32
native_by={s['objectId']:s for s in native['slides']}
slide_by={s['objectId']:p.slides[i] for i,s in enumerate(native['slides'])}
element_by={s['objectId']:{e['objectId']:list(p.slides[i].shapes)[j] for j,e in enumerate(s['pageElements'])} for i,s in enumerate(native['slides'])}
for s in native['slides']:
    assert len(s['pageElements'])==len(slide_by[s['objectId']].shapes)
NAVY=(32,54,92); DARK=(25,39,55); GREY=(92,104,119)

def write(shape,text,size,bold=False,color=DARK):
    tf=shape.text_frame
    tf.clear();tf.word_wrap=True
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(text.split('\n')):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.text=line;para.font.name='Arial';para.font.size=Pt(size)
        para.font.bold=bold;para.font.color.rgb=RGBColor(*color)
        para.space_before=Pt(0);para.space_after=Pt(0);para.line_spacing=1.06
        pr=para._p.get_or_add_pPr()
        for el in list(pr):
            if el.tag.rsplit('}',1)[-1] in ['buChar','buAutoNum','buBlip','buNone']:
                pr.remove(el)
        pr.append(OxmlElement('a:buNone'))
    return tf

def box(slide,text,x,y,width,height,size,bold=False,color=DARK):
    sh=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(width),Pt(height))
    sh.text_frame.margin_left=Pt(4);sh.text_frame.margin_right=Pt(4)
    sh.text_frame.margin_top=Pt(3);sh.text_frame.margin_bottom=Pt(3)
    write(sh,text,size,bold,color)
    return sh

source_media={}
for sid,slide in slide_by.items():
    source_media[sid]=[hashlib.sha256(sh.image.blob).hexdigest() for sh in slide.shapes if sh.shape_type==13]

content={s['id']:s for s in data['slides']}
content['p7']['title']='The Bayesian radius controls these bands'
for sid,s in content.items():
    sl=slide_by[sid]; es=element_by[sid]
    write(es[sid+'_i2'],s['title'],31.5,True,(255,255,255))
    write(es[sid+'_i3'],s['footer'],12.75,False,GREY)
    if s['type']=='story':
        for sh in list(sl.shapes):
            if sh.has_table:
                sh._element.getparent().remove(sh._element)
        for i,t in enumerate(s['body']):
            box(sl,t,47,112+i*102,866,91,21.75)
        box(sl,s['takeaway'],47,433,866,52,19.5,True,NAVY)
    else:
        for eid,t in s.get('texts',{}).items():
            tf=write(es[eid],t,19.5 if eid=='p3_i7' else 18,eid!='p3_i7',NAVY)
            if eid=='p3_i7':
                es[eid].top=Pt(155);es[eid].width=Pt(344);es[eid].height=Pt(286)
                for para in tf.paragraphs:para.space_after=Pt(14)
            if eid in ['p5_i7','analysis_coverage_plot_i7','analysis_halfwidth_plot_i7']:
                es[eid].top=Pt(445);es[eid].height=Pt(42)
            if eid=='p7_i11':
                es[eid].top=Pt(436);es[eid].height=Pt(28)
            if eid=='p7_i12':
                es[eid].top=Pt(467);es[eid].height=Pt(27)
                write(es[eid],t,16,False,NAVY)
    oldnotes=sl.notes_slide.notes_text_frame.text
    sl.notes_slide.notes_text_frame.text=(s.get('notes') or '\n\n'.join(s['body'])+'\n\n'+oldnotes)+'\n\nSource records: outputs/advisor_metrics/2026-09-29/METRIC_EVIDENCE.json; outputs/weighted_conformal/2026-09-29-v2/results.json; docs/WEIGHTED_CONFORMAL_DESIGN.md.\nThis update reuses saved evidence; it adds no training, simulation or synthetic trials.'

for i,sid in enumerate(data['order']):
    write(element_by[sid][sid+'_i4'],str(i+1),14)
    # Google export records table-frame sizes separately from native row/column sizes.
    # Align the frame to those preserved sizes for correct PowerPoint rendering.
    for sh in slide_by[sid].shapes:
        if sh.has_table:
            sh.width=sum(c.width for c in sh.table.columns)
            sh.height=sum(r.height for r in sh.table.rows)
    if i>=15 and sid not in content:
        title=element_by[sid][sid+'_i2']
        # Use notes, not oversized title prefixes, to identify the appendix.
        sl=slide_by[sid]
        sl.notes_slide.notes_text_frame.text='APPENDIX: supporting diagnostics and implementation details.\n\n'+sl.notes_slide.notes_text_frame.text

ids=list(p.slides._sldIdLst)
id_by={n['objectId']:ids[i] for i,n in enumerate(native['slides'])}
for el in ids:p.slides._sldIdLst.remove(el)
for sid in data['order']:p.slides._sldIdLst.append(id_by[sid])
for sid,el in id_by.items():
    if sid not in data['order']:p.part.drop_rel(el.rId)
dest=OUT/'BCA_Advisor_Update.pptx'
p.save(dest)
check=Presentation(dest)
assert len(check.slides)==26
for i,sid in enumerate(data['order']):
    images=[hashlib.sha256(sh.image.blob).hexdigest() for sh in check.slides[i].shapes if sh.shape_type==13]
    assert images==source_media[sid],sid
lines=['# BCA advisor update','', '15-slide presentation followed by 11 supporting slides.','']
for i,sid in enumerate(data['order']):
    sl=check.slides[i]
    title=next(sh.text for sh in sl.shapes if sh.has_text_frame and sh.top<Pt(50) and sh.text.strip())
    lines += [f'## {i+1}. {title}','',sl.notes_slide.notes_text_frame.text,'']
(OUT/'SPEAKER_NOTES.md').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'README.md').write_text('# BCA advisor update\n\nOpen `BCA_Advisor_Update.pptx` in PowerPoint. The first 15 slides tell the research story; the remaining 11 provide supporting diagnostics and implementation detail. Each result includes its purpose, interpretation and next test in the speaker notes.\n\nThis revision uses saved September 29 evidence. Original 32-slide files remain in `outputs/advisor_interpretation/2026-09-29/`. No scientific result, plot series or protocol changed.\n',encoding='utf-8')
(BASE/'powerpoint_build_receipt.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output':str(dest),'slides':26,'main_story':15,'appendix':11,'all_retained_plot_image_hashes_unchanged':True,'new_training':0,'new_simulator_steps':0,'new_synthetic_trials':0},indent=2))
print(json.dumps({'pptx':str(dest),'slides':len(check.slides),'images_preserved':True}))
